"""Builds ONSE deliverable F.3.1.1: curriculum syllabus, course outline, and implementation plan (Word).

Run: .venv/bin/python _build/build_curriculum_doc.py  ->  docs/ONSE-Business-Suite-Curriculum-and-Implementation-Plan.docx
"""
import re
from collections import Counter
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

import curriculum as C
from docx_kit import BLUE, INK, MUTED, RED, WHITE, Kit, borders, fix, shade

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "ONSE-Business-Suite-Curriculum-and-Implementation-Plan.docx"


def for_onse(text):
    """Internal 'CONFIRM' reminders become a statement of practice in the submitted document."""
    text = re.sub(r"\(CONFIRM current (.*?) before each cohort\)", r"(staff re-verify current \1 before each cohort)", text)
    return text.replace("(Proposed; confirm with ONSE.)", "(Proposed for ONSE approval.)")


def module_of(day):
    return next(m for m in C.MODULES if m["days"][0] <= day["n"] <= m["days"][1])


class Doc(Kit):
    def table(self, headers, widths, rows, size=9.5, bold_first=True):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        borders(t)
        for i, h in enumerate(headers):
            c = t.cell(0, i)
            shade(c, "3A56A6")
            r = c.paragraphs[0].add_run(h)
            r.bold, r.font.size, r.font.color.rgb = True, Pt(size), WHITE
        for ri, row in enumerate(rows, start=1):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                for li, line in enumerate(val if isinstance(val, list) else [val]):
                    p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    p.paragraph_format.space_after = Pt(1)
                    r = p.add_run(for_onse(str(line)))
                    r.font.size = Pt(size)
                    r.bold = bold_first and ci == 0
            if ri % 2 == 0:
                for ci in range(len(headers)):
                    shade(t.cell(ri, ci), "F6F8FC")
        fix(t, widths)
        self.para("", after=4)
        return t

    def h1(self, text, new_page=True):
        if new_page:
            self.doc.add_page_break()
        self.doc.add_heading(text, level=1)

    def h3(self, text):
        self.doc.add_heading(text, level=3)

    def body(self, text):
        return self.para(for_onse(text))

    def labeled(self, label, text, after=2):
        p = self.para(label, bold=True, after=after)
        p.add_run(for_onse(text)).bold = False

    def bullet_list(self, lines):
        for line in lines:
            p = self.doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            if isinstance(line, tuple):
                p.add_run(line[0] + ". ").bold = True
                p.add_run(for_onse(line[1]))
            else:
                p.add_run(for_onse(line))

    def page_numbers(self):
        p = self.sec.footer.paragraphs[0]
        p.text = f"{C.PROGRAM['title']} · ONSE Contract {C.PROGRAM['contract']} · The Chart Academy    Page "
        for kind in ("begin", "instr", "end"):
            r = p.add_run()
            if kind == "instr":
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = "PAGE"
            else:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(8.5), MUTED


def block_kind(label):
    return {
        "Expert Content": "Expert Content (instruction)",
        "Master Class": "Master Classes (guest experts and Truist)",
        "Field Experience": "Field Experience (applied, real-world)",
        "Build": "Portfolio build and capstone",
        "Capstone": "Portfolio build and capstone",
    }.get(label, "Orientation, check-in, tech catch-up, and assessment")


def cover(d):
    doc = d.doc
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(50)
    p.add_run().add_picture(str(ROOT / "images" / "chart-logo-full.png"), width=Inches(2.0))
    for text, size, color, bold in (
        ("Office of Neighborhood Safety and Engagement · Pathways Program", 11, MUTED, True),
        (C.PROGRAM["title"], 30, INK, True),
        (C.PROGRAM["subtitle"], 16, RED, True),
        ("Curriculum Syllabus, Course Outline, and Implementation Plan", 14, BLUE, True),
        (f"Deliverable F.3.1.1 · Contract {C.PROGRAM['contract']} · Solicitation {C.PROGRAM['solicitation']}", 10.5, MUTED, False),
    ):
        q = doc.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        q.paragraph_format.space_after = Pt(8)
        r = q.add_run(text)
        r.font.size, r.font.color.rgb, r.bold = Pt(size), color, bold
    q = doc.add_paragraph()
    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q.paragraph_format.space_before = Pt(50)
    for line in ("Submitted to: ONSE Contract Administrator",
                 "Submitted by: " + C.PROGRAM["provider"],
                 "529 14th Street NW, Suite 718, Washington, DC 20045 · 202.989.5621",
                 "Point of contact: Cedric Thompson, CEO · cedric.thompson@we-strategies.us",
                 "Curriculum partner: " + C.PROGRAM["partner"],
                 "Version 1.0 · October 2026"):
        r = q.add_run(line)
        r.font.size = Pt(10)
        r.add_break()


def contents(d):
    d.h1("Contents")
    titles = ["Program Overview", "Requirements Crosswalk", "Learning Objectives", "Scope and Sequence", "Time Allocation",
              "Daily Lesson Plans", "Instructional Materials", "Teaching Strategies", "Assessment Plan",
              "Standards Alignment", "Support for Diverse Learners and Reentry", "Program Enhancements",
              "Evaluation and Feedback", "Implementation Plan"]
    d.table(["Section", "Title"], [1.0, 5.7], [(str(i), t) for i, t in enumerate(titles, 1)], size=10.5)


def overview(d):
    P = C.PROGRAM
    d.h1("1. Program Overview")
    d.body(f"The {P['title']} is a pre-existing, ready-to-implement entrepreneurship program delivered in person to ONSE "
           "Pathways cohorts. It is built on the Institute of Research for Social Justice and Action (IRSJA) curriculum "
           "(Boss Up, Money Moves, and Flipping the Game), which has been implemented nationally with justice-impacted "
           "adults, serves more than 800 participants a year, and is approved for use by the Federal Bureau of Prisons. "
           "The Chart Academy delivers it through its in-person instructional model, Canvas learning management system, "
           "case management, and District employer and financial partnerships.")
    d.labeled("Who it serves. ", P["population"] + " The curriculum is designed for that reality: a predictable daily "
              "routine, plain-language materials, a short daily Tech Catch-Up, a readiness review in week one, and a "
              "dedicated Case Management Associate.")
    d.labeled("What participants leave with. ", "A launch-ready business portfolio (business plan and DC compliance pathway, "
              "brand kit, financial plan, marketing strategy, and 90-day launch plan), a resume, IRSJA reference guides to "
              "keep, continued Canvas access, a one-year Greater Washington Black Chamber of Commerce membership, a banking "
              "relationship through Truist Bank, and a professional network built through Master Classes and the pitch panel.",
              after=8)
    d.table(["Item", "Detail"], [2.0, 4.7], [
        ("Contract", f"{P['contract']} (solicitation {P['solicitation']}), {P['agency']}"),
        ("Instruction", f"{P['hours']} in-person hours over {P['days']} instructional days ({P['hours_per_day']} hours a day)"),
        ("Cohort size", f"{P['cohort_size']}; up to 3 cohorts per contract year"),
        ("Schedule", P["schedule"]),
        ("Location", P["location"]),
        ("Structure", "Four modules: BUILD (Days 1 to 5), BRAND (Days 6 to 10), FINANCE (Days 11 to 15), LAUNCH (Days 16 to 20)"),
        ("Capstone", "Business portfolio and a pitch competition before an external panel on Day 20"),
        ("Platform", "Canvas LMS (thechartacademy.instructure.com) and the course website"),
    ])


def crosswalk(d):
    d.h1("2. Requirements Crosswalk")
    d.body("Where each requirement in Section C.5, and each curriculum element listed in C.5.2, is addressed.")
    d.table(["RFP", "Requirement", "Where it is met"], [0.8, 3.4, 2.5], [
        ("C.5.1", "ONSE-approved, pre-existing curriculum implemented with adult learners", "Section 1"),
        ("C.5.2", "Integrated curriculum: instruction, workbooks, capstone, enhancements, reporting", "Sections 3 to 14"),
        ("C.5.2.1", "Program enhancements: social capital, mentorship, financial institution", "Section 12"),
        ("C.5.3", "60 in-person instructional hours over 15 to 20 days", "Sections 4, 5, 6"),
        ("C.5.4", "Cohorts of 15 to 20 at a District-approved location", "Sections 1, 14"),
        ("C.5.5", "Required topics, including resumes and interviews", "Sections 5, 6, 10"),
        ("C.5.6", "Business 101 Reference Guide or equivalent workbook", "Section 7"),
        ("C.5.7", "Capstone activity at the end of each cohort", "Sections 6 (Day 20), 9"),
        ("C.5.8", "Pre- and post-assessments; attendance, participation, and outcomes", "Sections 9, 13, 14"),
        ("C.5.9", "Social capital development activities", "Section 12"),
        ("C.5.10", "Mentorship and entrepreneurial guidance", "Section 12"),
        ("C.5.11", "Financial institution participation", "Section 12"),
        ("C.5.12", "Lead Instructor and SME qualifications", "Section 14"),
    ])
    d.table(["C.5.2 curriculum element", "Section"], [4.7, 2.0], [
        ("Learning objectives", "3"), ("Content scope and sequence", "4"), ("Instructional materials", "7"),
        ("Teaching strategies", "8"), ("Assessment plan", "9"), ("Standards alignment", "10"),
        ("Time allocation", "5"), ("Support for diverse learners", "11"), ("Evaluation and feedback", "13"),
    ])


def objectives(d):
    d.h1("3. Learning Objectives")
    d.doc.add_heading("Program objectives", level=2)
    d.body("By the end of the program, each participant will be able to:")
    d.bullet_list([
        "Choose a legal business structure and complete a DC compliance pathway (registration, EIN, Clean Hands, licenses).",
        "Validate demand through customer interviews and survey data, and build a brand on that evidence.",
        "Market the business through social media and a content strategy, and protect its name and ideas.",
        "Set prices, build a revenue model, and create startup and monthly budgets.",
        "Use business banking, understand credit, and identify funding sources.",
        "Communicate professionally: elevator pitch, networking, resume, interview, and a formal stakeholder pitch.",
        "Set personal financial goals and shift from short-term survival decisions to long-term planning.",
    ])
    for m in C.MODULES:
        d.doc.add_heading(f"Module {m['num']}: {m['name']} · {m['focus']}", level=2)
        d.bullet_list(m["objectives"])


def scope_sequence(d):
    d.h1("4. Scope and Sequence")
    d.body("Each module builds on the one before it, moving from setup to brand to money to launch. Every day follows the "
           "same rhythm, so participants always know what comes next:")
    d.table(["Part", "Time", "What happens"], [1.7, 1.1, 3.9], C.DAY_RHYTHM)
    d.table(["Module", "Days", "Reference guide", "Module project", "Outcome"], [1.1, 0.7, 1.1, 2.0, 1.8], [
        (f"{m['num']}. {m['name']}", f"{m['days'][0]} to {m['days'][1]}", m["guide"], m["project"], m["outcome"])
        for m in C.MODULES])
    d.doc.add_heading("20-day sequence", level=2)
    rows = []
    for day in C.DAYS:
        extra = [x for x in (day.get("master") and "Master Class: " + day["master"],
                             day.get("truist") and "Truist: " + day["truist"],
                             day.get("field") and "Field: " + day["field"]) if x]
        rows.append((str(day["n"]), module_of(day)["name"], day["title"],
                     [C.TOPICS[t] for t in day["topics"]], extra or ["Classroom lesson and portfolio build"]))
    d.table(["Day", "Module", "Lesson", "Required topics", "Field, Master Class, or partner"],
            [0.45, 0.8, 1.6, 1.85, 2.0], rows, size=8)


def time_allocation(d):
    d.h1("5. Time Allocation")
    by_kind, by_mod, topic_min = Counter(), Counter(), Counter()
    for day in C.DAYS:
        for mins, label, _ in day["agenda"]:
            by_kind[block_kind(label)] += mins
            by_mod[module_of(day)["name"]] += mins
        taught = sum(m for m, label, _ in day["agenda"] if label in ("Expert Content", "Master Class", "Field Experience"))
        for t in day["topics"]:
            topic_min[t] += taught / len(day["topics"])
    total = sum(by_mod.values())
    d.table(["Module", "Days", "Hours"], [3.0, 2.0, 1.7],
            [(f"{m['num']}. {m['name']}", f"{m['days'][0]} to {m['days'][1]}", f"{by_mod[m['name']] / 60:g}") for m in C.MODULES]
            + [("Total", "20", f"{total / 60:g}")])
    d.table(["Instructional component", "Hours", "Share"], [3.9, 1.4, 1.4],
            [(k, f"{v / 60:.1f}", f"{v / total:.0%}") for k, v in by_kind.most_common()])
    d.doc.add_heading("Hours by required topic (C.5.5)", level=2)
    d.body("Counted from the Expert Content, Master Class, and Field Experience blocks of the days that teach each topic, "
           "split evenly when a day teaches more than one. Topics are also practiced in the daily portfolio build.")
    d.table(["Topic", "Days taught", "Hours"], [3.8, 1.9, 1.0],
            [(C.TOPICS[t], ", ".join(str(x["n"]) for x in C.DAYS if t in x["topics"]), f"{topic_min[t] / 60:.1f}")
             for t in C.TOPICS])


def lessons(d):
    d.h1("6. Daily Lesson Plans")
    d.body("Each lesson lists its objectives, a minute-by-minute agenda (180 minutes), the exit ticket, and the reentry "
           "support built into that day. Facilitators adapt pacing to the cohort while keeping every required block.")
    for m in C.MODULES:
        d.doc.add_page_break()
        d.doc.add_heading(f"Module {m['num']}: {m['name']} · Days {m['days'][0]} to {m['days'][1]}", level=2)
        d.labeled("Reference guide: ", m["guide"])
        d.body(m["summary"])
        d.para(f"Module project: {m['project']}", bold=True, after=2)
        d.bullet_list(m["project_items"])
        for day in C.module_days(m):
            d.h3(f"Day {day['n']}: {day['title']}")
            d.para("Objectives", bold=True, after=1)
            d.bullet_list(day["objectives"])
            d.table(["Min", "Block", "Activity"], [0.5, 1.3, 4.9],
                    [(str(mins), label, text) for mins, label, text in day["agenda"]], size=9)
            if day.get("exit"):
                d.labeled("Exit ticket: ", day["exit"])
            if day.get("reentry"):
                d.labeled("Reentry support: ", day["reentry"], after=8)


def materials(d):
    d.h1("7. Instructional Materials")
    d.table(["Material", "Description"], [2.1, 4.6], [
        ("IRSJA reference guides (C.5.6)", "Boss Up (Modules 1 and 2), Money Moves (Module 3), and Flipping the Game "
         "(Module 4 and the behavioral framework throughout). Each participant receives printed guides by Day 1 and keeps "
         "them. Together they serve as the Business 101 Reference Guide and cover every topic in C.5.6.1 (a to j)."),
        ("Participant workbook", "Fill-in worksheets for each day's Build block, the module project templates, and the "
         "capstone portfolio checklist."),
        ("Course website", "Course site with each day's objectives, agenda, links, and reentry resources."),
        ("Canvas LMS", "Daily exit tickets (attendance and understanding), assignment submission, rubrics, templates, the "
         "pre- and post-assessment, and continued access after graduation."),
        ("Templates", "One-page business plan, DC compliance checklist, interview script, survey starter, brand kit, 30-day "
         "content plan, revenue model and budget spreadsheet, capability statement, resume, and pitch deck."),
        ("Technology", "A laptop and professional bag for each participant; Google Workspace, Canva, and an AI assistant."),
        ("Partner materials", "Truist Bank Bank On Your Success session materials; DLCP, DSLBD, and SCORE handouts."),
    ])
    d.doc.add_heading("Business 101 Reference Guide topics (C.5.6.1)", level=2)
    d.table(["Topic", "Guide", "Days"], [3.4, 2.2, 1.1], [
        ("a. Business finance and financial literacy", "Money Moves", "11, 12"),
        ("b. Pricing strategies and revenue models", "Money Moves", "13"),
        ("c. Branding and marketing fundamentals", "Boss Up", "7, 8"),
        ("d. Social media marketing strategies", "Boss Up", "9"),
        ("e. Intellectual property and trademark basics", "Boss Up", "5, 8"),
        ("f. Licensing and regulatory compliance", "Boss Up", "3, 4"),
        ("g. Business banking and financial services", "Money Moves; Truist", "3, 8, 14, 18"),
        ("h. Public speaking and professional communication", "Boss Up; Flipping the Game", "10, 18"),
        ("i. Networking and relationship building", "Boss Up", "10"),
        ("j. Wealth mindset and goal setting", "Money Moves; Flipping the Game", "1, 11, 15"),
    ])


def teaching(d):
    d.h1("8. Teaching Strategies")
    d.bullet_list(C.TEACHING)
    d.doc.add_heading("Behavioral transformation: Flipping the Game", level=2)
    d.body("Through structured reflection and facilitated discussion threaded through all four modules, participants shift "
           "from short-term to long-term planning; build discipline, accountability, and goal-setting habits; strengthen "
           "emotional awareness and decision-making; and move from informal or high-risk economic activity to structured "
           "business pathways.")


def assessment(d):
    d.h1("9. Assessment Plan")
    d.table(["Assessment", "When", "What"], [1.6, 1.4, 3.7], C.ASSESSMENT)
    d.h3("Course grade")
    d.table(["Component", "What counts", "Weight"], [1.6, 4.1, 1.0],
            [(n, desc, f"{w}%") for n, desc, w in C.WEIGHTS] + [("Total", "", f"{sum(w for _, _, w in C.WEIGHTS)}%")])
    d.labeled("Completion. ", C.COMPLETION, after=8)
    for m in C.MODULES:
        if m["rubric"]:
            d.h3(f"Module {m['num']} project rubric: {m['project']} (20 points)")
            d.table(["Criterion", "Full · 5", "Partial · 3", "Not yet · 1"], [1.4, 2.1, 1.8, 1.4], m["rubric"], size=8.5)
    d.h3("Capstone rubric (50 points)")
    d.body(C.CAPSTONE["summary"])
    d.table(["Criterion", "Points", "What the panel looks for"], [1.6, 0.8, 4.3],
            [(name, str(pts), desc) for name, pts, desc in C.CAPSTONE["rubric"]])
    d.body(C.CAPSTONE["awards"])


def standards(d):
    d.h1("10. Standards Alignment")
    d.table(["Standard", "Alignment"], [2.6, 4.1], C.STANDARDS)
    d.doc.add_heading("Required topic crosswalk (C.5.5)", level=2)
    assessed = {"FIN": "Module 3 project; pre/post", "PRICE": "Module 3 project; capstone",
                "BRAND": "Module 2 project; capstone", "SOCIAL": "Module 2 project", "IP": "Module 1 and 2 projects",
                "LIC": "Module 1 project", "SPEAK": "Elevator pitch; capstone pitch", "NET": "Networking follow-up; capstone",
                "WEALTH": "Module 3 goal statement", "RESUME": "Resume in portfolio; mock interview",
                "BANK": "Module 3 project; pre/post"}
    d.table(["Required topic", "Days", "Assessed in"], [3.0, 1.4, 2.3],
            [(C.TOPICS[t], ", ".join(str(x["n"]) for x in C.DAYS if t in x["topics"]), assessed[t]) for t in C.TOPICS])


def diverse(d):
    d.h1("11. Support for Diverse Learners and Reentry")
    d.body("Many participants are coming home after months or years of incarceration. They may be re-entering a city, an "
           "economy, and a technology landscape that changed while they were gone, while also managing supervision, housing, "
           "identification, debt, and family. The curriculum treats these as design requirements, not side issues.")
    d.bullet_list(C.DIVERSE_LEARNERS)
    d.doc.add_heading("Reentry supports built into the lessons", level=2)
    d.table(["Day", "Reentry support"], [0.6, 6.1], [(str(x["n"]), x["reentry"]) for x in C.DAYS if x.get("reentry")], size=9)
    d.doc.add_heading("Tech Catch-Up sequence", level=2)
    d.bullet_list(C.TECH_CATCHUP)


def enhancements(d):
    d.h1("12. Program Enhancements")
    d.table(["Enhancement", "How it is delivered"], [2.1, 4.6], C.ENHANCEMENTS)
    d.doc.add_heading("Truist Bank: Bank On Your Success topics", level=2)
    d.bullet_list(["Business finance basics", "Business structure and organizing your finances",
                   "Separating business and personal finances", "Financial accounts and tools", "Managing business cash flow",
                   "How to pay yourself", "Borrowing for small business; where to find funding",
                   "What lenders look for; creating your borrowing profile", "Building, expanding, and protecting your business"])


def evaluation(d):
    d.h1("13. Evaluation and Feedback")
    d.bullet_list([
        ("Participant growth", "Pre- and post-assessment results compared for each participant and the cohort."),
        ("Daily feedback", "Exit tickets show what each participant understood; facilitators adjust the next day."),
        ("Assignment feedback", "Written rubric feedback in Canvas on each module project within two instructional days."),
        ("Panel feedback", "Pitch panelists give each participant written scores and comments."),
        ("Program evaluation", "End-of-program participant survey on instruction, supports, and enhancements."),
        ("ONSE feedback", "Debrief with ONSE after each cohort; changes are documented before the next cohort."),
        ("Continuous improvement", "The Program Director and curriculum developer review results after each cohort and "
                                   "update pacing, examples, and resources while keeping curriculum fidelity."),
    ])


def implementation(d):
    d.h1("14. Implementation Plan")
    d.doc.add_heading("Staffing", level=2)
    d.table(["Role", "Name", "Responsibilities"], [1.9, 1.6, 3.2], C.STAFF)
    d.body("The Lead Instructor has more than three years of experience delivering business, entrepreneurship, and workforce "
           "training to adult learners, and the Subject Matter Expert is the developer of the curriculum (C.5.12).")
    d.doc.add_heading("Cohort timeline", level=2)
    d.table(["When", "Activity", "Lead"], [1.7, 3.5, 1.5], [
        ("10 business days before Day 1", "Confirm schedule, site, and roster with ONSE; set up the Canvas cohort.",
         "Program Director"),
        ("5 business days before Day 1", "Confirm Master Class speakers, Truist sessions, SCORE mentors, and the pitch panel; "
                                         "print reference guides and workbooks; prepare laptops.", "Lead Instructor, Operations"),
        ("Day 1", "Orientation, pre-assessment, Canvas setup; pre-assessment summary to ONSE.", "Lead Instructor, Data Analyst"),
        ("Week 1", "Readiness review with every participant.", "Case Management Associate"),
        ("Days 1 to 20", "Daily instruction, attendance, and progress reports.", "Lead Instructor, Data Analyst"),
        ("Days 5, 10, 15", "Module projects scored with feedback.", "Lead Instructor"),
        ("Day 20", "Capstone pitch competition, post-assessment, and graduation.", "All staff"),
        ("Within 15 business days of the capstone", "Final cohort report to ONSE.", "Data Analyst, Program Director"),
    ])
    d.doc.add_heading("Deliverables (Section F.3.1)", level=2)
    d.table(["#", "Deliverable", "Format", "Due"], [0.75, 3.0, 1.3, 1.65], [
        ("F.3.1.1", "Curriculum syllabus, course outline, and implementation plan (this document)", "Electronic",
         "Within 10 business days after award"),
        ("F.3.1.2", "Participant reference guides and workbooks (15 to 20)", "Hard copy and electronic",
         "By the first instructional day"),
        ("F.3.1.3", "Pre-assessment tool and results summary", "Electronic", "On or before Day 1"),
        ("F.3.1.4", "In-person instruction: 60 hours over 20 days", "In person", "ONSE-approved schedule"),
        ("F.3.1.5", "Capstone activity (pitch competition)", "In person, with electronic artifacts", "Day 20"),
        ("F.3.1.6", "Progress and attendance reports (one per participant)", "Electronic", "Daily"),
        ("F.3.1.7", "Final cohort report: attendance, completion, outcomes, assessment results", "Electronic",
         "Within 15 business days after the capstone"),
    ], size=9)
    d.doc.add_heading("Attendance, case management, and reporting", level=2)
    d.bullet_list([
        "Attendance is recorded daily through the Canvas exit ticket and the sign-in sheet, and reviewed the same day.",
        "Two or more consecutive absences trigger outreach from the Case Management Associate within one business day.",
        "Barriers are logged with the support provided (transportation, supervision conflicts, housing, ID, childcare).",
        "Daily progress and attendance reports go to ONSE, and a live cohort dashboard shows attendance, assignment "
        "completion, and assessment results.",
        "The final cohort report includes attendance and completion rates, pre/post growth, capstone results, enhancement "
        "participation (GWBCC memberships, Truist accounts opened, mentor matches), and participant feedback.",
    ])


def main():
    C.check()
    d = Doc(str(ROOT / "images" / "chart-logo-icon.png"))
    d.page_numbers()
    cover(d)
    for part in (contents, overview, crosswalk, objectives, scope_sequence, time_allocation, lessons, materials, teaching,
                 assessment, standards, diverse, enhancements, evaluation, implementation):
        part(d)
    OUT.parent.mkdir(exist_ok=True)
    d.save(str(OUT))
    print("Wrote", OUT.relative_to(ROOT))


if __name__ == "__main__":
    main()
